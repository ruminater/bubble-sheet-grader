"""Results as files: CSV, an Excel workbook, a class report PDF, student pages PDF.

All four are built from the same Grading, so they always agree with the page.
Rosters are alphabetical by name (the name grid is LAST FIRST), unnamed sheets last.
"""

from __future__ import annotations

import csv
import io
import re

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .grade import Grading, Result

# One palette for every output. Purple is the web page's accent.
PURPLE = "4F2D7F"
PURPLE_LIGHT = "EDE7F6"
INK = "1D2330"
MUTED = "5D6675"
LINE = "DFE3EA"
ZEBRA = "F6F7F9"
GOOD, GOOD_BG = "167A3E", "E3F4E8"
BAD, BAD_BG = "B3261E", "FBE4E2"
WARN, WARN_BG = "A15C00", "FFF1D6"

CHOICES = "ABCDE"


def filename(g: Grading, ext: str, suffix: str = "") -> str:
    base = " ".join(x for x in (g.course, g.title) if x) or "bubble sheet results"
    base = re.sub(r"[^A-Za-z0-9 _.-]+", "", base).strip() or "bubble sheet results"
    return f"{base}{suffix} {g.graded_at:%Y-%m-%d}.{ext}"


def _day(d, month: str = "%B") -> str:
    # Built by hand: strftime's %-d (no leading zero) is not supported on Windows.
    return f"{d.strftime(month)} {d.day}, {d.year}"


def _stamp(g: Grading) -> str:
    d = g.graded_at
    return f"{_day(d)} at {d.hour % 12 or 12}:{d:%M %p}"


def _response(r: Result, q: int) -> str:
    a = r.responses[q]
    return {"": "-", "*": "Multiple"}.get(a, a)


# --------------------------------------------------------------------- CSV


def to_csv(g: Grading) -> str:
    """One row per student, key row first. BOM so Excel opens it as UTF-8."""
    buf = io.StringIO()
    buf.write("﻿")
    w = csv.writer(buf)
    qs = g.questions
    w.writerow(["Name", "Sheet", "Score", "Possible", "Percent", "Marks to check"] + [f"Q{q}" for q in qs])
    w.writerow(["ANSWER KEY", "", "", "", "", ""] + [g.key_answers[q] for q in qs])
    for r in g.by_name:
        w.writerow(
            [r.sheet.name or f"(no name, sheet {r.sheet.number})", r.sheet.number, r.score,
             r.possible, f"{r.percent:.1f}", len(r.flags)]
            + [{"": "", "*": "*"}.get(r.responses[q], r.responses[q]) for q in qs]
        )
    return buf.getvalue()


# ------------------------------------------------------------------- Excel

_thin = Side(style="thin", color=LINE)
_BORDER = Border(bottom=_thin)
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEAD_FILL = PatternFill("solid", fgColor=PURPLE)
_ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")


def _fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def _title(ws, g: Grading, subtitle: str):
    ws["A1"] = g.heading
    ws["A1"].font = Font(bold=True, size=18, color=PURPLE)
    ws["A2"] = " · ".join(x for x in (g.course, subtitle, f"Graded {_stamp(g)}") if x)
    ws["A2"].font = Font(size=10, color=MUTED)
    ws.row_dimensions[1].height = 28


def _header(ws, row: int, labels: list[str], widths: list[float] | None = None):
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font, cell.fill = _HEAD_FONT, _HEAD_FILL
        cell.alignment = _CENTER if c > 1 else _LEFT
    ws.row_dimensions[row].height = 22
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def _zebra(ws, first: int, last: int, cols: int):
    for r in range(first, last + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _BORDER
            if (r - first) % 2 == 1 and cell.fill.fgColor.rgb in (None, "00000000"):
                cell.fill = _ZEBRA_FILL


def _print_setup(ws, header_row: int, landscape: bool = False):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.sheet_view.showGridLines = False


def to_xlsx(g: Grading) -> bytes:
    wb = Workbook()
    _xlsx_summary(wb.active, g)
    _xlsx_scores(wb.create_sheet("Scores"), g)
    _xlsx_responses(wb.create_sheet("Responses"), g)
    _xlsx_items(wb.create_sheet("Item Analysis"), g)
    _xlsx_checks(wb.create_sheet("To Check"), g)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _xlsx_summary(ws, g: Grading):
    ws.title = "Summary"
    _title(ws, g, "Summary")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    s = g.stats()
    rows = [
        ("Students", s.get("count", 0), "0"),
        ("Questions", len(g.questions), "0"),
        ("Mean", s.get("mean", 0) / 100, "0.0%"),
        ("Median", s.get("median", 0) / 100, "0.0%"),
        ("Highest", s.get("high", 0) / 100, "0.0%"),
        ("Lowest", s.get("low", 0) / 100, "0.0%"),
        ("Standard deviation", s.get("stdev", 0) / 100, "0.0%"),
    ]
    ws["A4"] = "At a glance"
    ws["A4"].font = Font(bold=True, size=12, color=INK)
    for i, (label, value, fmt) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=label).font = Font(color=MUTED)
        c = ws.cell(row=i, column=2, value=value)
        c.number_format, c.font, c.alignment = fmt, Font(bold=True, size=12), Alignment(horizontal="right")
        ws.cell(row=i, column=1).border = ws.cell(row=i, column=2).border = _BORDER

    # Score distribution, with a chart beside it.
    top = 5 + len(rows) + 2
    ws.cell(row=top, column=1, value="Score distribution").font = Font(bold=True, size=12)
    _header(ws, top + 1, ["Percent", "Students"])
    dist = g.distribution()
    first, last = top + 2, top + 1 + len(dist)
    for i, (label, n) in enumerate(dist, first):
        ws.cell(row=i, column=1, value=label + "%")
        c = ws.cell(row=i, column=2, value=n)
        c.alignment = _CENTER
    _zebra(ws, first, last, 2)
    chart = BarChart()
    chart.type, chart.style = "col", 10
    chart.title, chart.legend = "Students by score", None
    chart.y_axis.title, chart.x_axis.title = "Students", "Percent correct"
    chart.y_axis.majorGridlines = None
    chart.y_axis.delete = chart.x_axis.delete = False
    chart.height, chart.width = 7.5, 15
    chart.y_axis.majorUnit = max(1, round(max(n for _, n in dist) / 4))
    chart.y_axis.number_format = "0"
    chart.add_data(Reference(ws, min_col=2, min_row=first, max_row=last), titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    chart.series[0].graphicalProperties.solidFill = PURPLE
    chart.series[0].graphicalProperties.line.solidFill = PURPLE
    ws.add_chart(chart, "D4")

    notes_row = top + 3 + len(dist)
    notes = g.key_warnings + g.skipped
    if notes:
        ws.cell(row=notes_row, column=1, value="Notes").font = Font(bold=True, size=12)
        for i, n in enumerate(notes, notes_row + 1):
            c = ws.cell(row=i, column=1, value=n)
            c.font, c.alignment = Font(color=WARN), Alignment(wrap_text=False)
        notes_row += len(notes) + 2
    if g.sources:
        ws.cell(row=notes_row, column=1, value="Files graded").font = Font(bold=True, size=12)
        for i, n in enumerate(g.sources, notes_row + 1):
            ws.cell(row=i, column=1, value=n).font = Font(color=MUTED)
    _print_setup(ws, 1)


def _xlsx_scores(ws, g: Grading):
    _title(ws, g, "Scores")
    _header(ws, 4, ["Name", "Sheet", "Score", "Possible", "Percent", "Marks to check"],
            [30, 8, 9, 10, 12, 16])
    rows = g.by_name
    for i, r in enumerate(rows, 5):
        vals = [r.sheet.name or f"(no name, sheet {r.sheet.number})", r.sheet.number, r.score,
                r.possible, r.percent / 100, len(r.flags) or None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.alignment = _LEFT if c == 1 else _CENTER
        ws.cell(row=i, column=5).number_format = "0.0%"
        if r.flags:
            ws.cell(row=i, column=6).font = Font(bold=True, color=WARN)
        if not r.sheet.name:
            ws.cell(row=i, column=1).font = Font(italic=True, color=MUTED)
    last = 4 + len(rows)
    _zebra(ws, 5, last, 6)
    if rows:
        ws.conditional_formatting.add(
            f"E5:E{last}",
            DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="B39DDB"),
        )
        ws.auto_filter.ref = f"A4:F{last}"
    ws.freeze_panes = "B5"
    _print_setup(ws, 4)


def _xlsx_responses(ws, g: Grading):
    _title(ws, g, "Responses: green correct, red wrong, amber to check, - blank, * two bubbles")
    qs = g.questions
    _header(ws, 4, ["Name", "Score"] + [str(q) for q in qs], [28, 8] + [4.2] * len(qs))
    key_row = 5
    ws.cell(row=key_row, column=1, value="ANSWER KEY").font = Font(bold=True, color=PURPLE)
    for j, q in enumerate(qs, 3):
        c = ws.cell(row=key_row, column=j, value=g.key_answers[q])
        c.font, c.alignment, c.fill = Font(bold=True, color=PURPLE), _CENTER, _fill(PURPLE_LIGHT)
    ws.cell(row=key_row, column=1).fill = ws.cell(row=key_row, column=2).fill = _fill(PURPLE_LIGHT)
    good, bad, warn = _fill(GOOD_BG), _fill(BAD_BG), _fill(WARN_BG)
    for i, r in enumerate(g.by_name, key_row + 1):
        ws.cell(row=i, column=1, value=r.sheet.name or f"(no name, sheet {r.sheet.number})")
        sc = ws.cell(row=i, column=2, value=f"{r.score}/{r.possible}")
        sc.alignment = _CENTER
        for j, q in enumerate(qs, 3):
            a = r.responses[q]
            c = ws.cell(row=i, column=j, value=a or "-")
            c.alignment = _CENTER
            if q in r.flags:
                c.fill, c.font = warn, Font(bold=True, color=WARN)
            elif q in r.correct:
                c.fill, c.font = good, Font(color=GOOD)
            else:
                c.fill, c.font = bad, Font(bold=True, color=BAD)
            c.border = Border(left=_thin, right=_thin, bottom=_thin)
    ws.freeze_panes = "C6"
    ws.print_title_cols = "A:B"
    _print_setup(ws, 4, landscape=True)


def _xlsx_items(ws, g: Grading):
    _title(ws, g, "Item analysis")
    labels = ["Question", "Key", "% Correct"] + list(CHOICES) + ["Blank", "Multiple", "Discrimination", "Worth a look"]
    _header(ws, 4, labels, [10, 7, 11, 6, 6, 6, 6, 6, 8, 10, 15, 46])
    for i, it in enumerate(g.items, 5):
        vals = [it.question, it.key, it.percent / 100] + [it.counts[c] or None for c in CHOICES] + [
            it.counts[""] or None, it.counts["*"] or None, it.discrimination, it.note or None]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.alignment = _LEFT if c == len(vals) else _CENTER
        ws.cell(row=i, column=3).number_format = "0%"
        ws.cell(row=i, column=11).number_format = "0.00"
        key_col = 4 + CHOICES.index(it.key)
        ws.cell(row=i, column=key_col).font = Font(bold=True, color=GOOD)
        ws.cell(row=i, column=2).font = Font(bold=True, color=PURPLE)
        if it.note:
            ws.cell(row=i, column=12).font = Font(color=BAD if it.discrimination is not None and it.discrimination < 0 else WARN)
    last = 4 + len(g.items)
    _zebra(ws, 5, last, len(labels))
    if g.items:
        ws.conditional_formatting.add(
            f"C5:C{last}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F4B6B0",
                           mid_type="num", mid_value=0.6, mid_color="FFF1D6",
                           end_type="num", end_value=1, end_color="BFE5C9"),
        )
        ws.auto_filter.ref = f"A4:L{last}"
    ws.freeze_panes = "B5"
    r = last + 2
    for line in (
        "% Correct: share of students who chose the key.",
        "Discrimination: share correct in the top third by score minus the bottom third. "
        "0.3 or more is good; below 0 usually means a wrong key or a misleading question.",
    ):
        ws.cell(row=r, column=1, value=line).font = Font(italic=True, color=MUTED, size=9)
        r += 1
    _print_setup(ws, 4, landscape=True)


def _xlsx_checks(ws, g: Grading):
    _title(ws, g, "Marks a person should confirm on the paper sheet")
    _header(ws, 4, ["Name", "Sheet", "Question", "Issue", "Read as", "Key", "What it means"],
            [28, 8, 10, 11, 12, 7, 60])
    meaning = {
        "multiple": "Two bubbles filled. No credit given, as a machine would score it.",
        "light": "Faint or an X. Scored as read; confirm on the paper.",
    }
    rows = [(r, q, f) for r in g.by_name for q, f in sorted(r.flags.items())]
    for i, (r, q, f) in enumerate(rows, 5):
        vals = [r.sheet.name or f"(no name, sheet {r.sheet.number})", r.sheet.number, q,
                f.capitalize(), _response(r, q), g.key_answers[q], meaning[f]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.alignment = _LEFT if c in (1, 7) else _CENTER
        ws.cell(row=i, column=4).font = Font(bold=True, color=WARN)
    if not rows:
        ws.cell(row=5, column=1, value="Nothing to check: every mark was read cleanly.").font = Font(color=GOOD)
    _zebra(ws, 5, 4 + len(rows), 7)
    ws.freeze_panes = "A5"
    _print_setup(ws, 4, landscape=True)


# --------------------------------------------------------------------- PDF


def _c(hex_):
    return colors.HexColor("#" + hex_)


_BASE = ParagraphStyle("base", fontName="Helvetica", fontSize=9.5, leading=12.5, textColor=_c(INK))
_SMALL = ParagraphStyle("small", parent=_BASE, fontSize=8, leading=10, textColor=_c(MUTED))
_H1 = ParagraphStyle("h1", parent=_BASE, fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=_c(PURPLE))
_H2 = ParagraphStyle("h2", parent=_BASE, fontName="Helvetica-Bold", fontSize=13, leading=16,
                     spaceBefore=14, spaceAfter=6, textColor=_c(INK))
_SUB = ParagraphStyle("sub", parent=_BASE, fontSize=10, textColor=_c(MUTED), spaceAfter=10)
_CELL = ParagraphStyle("cell", parent=_BASE, fontSize=8.5, leading=10.5)
_TILE_NUM = ParagraphStyle("tn", parent=_BASE, fontName="Helvetica-Bold", fontSize=17, leading=20, alignment=TA_CENTER)
_TILE_LBL = ParagraphStyle("tl", parent=_SMALL, alignment=TA_CENTER)


class _NumberedCanvas(rl_canvas.Canvas):
    """Draws 'Page n of m' and the running header once the page count is known."""

    header = ""
    count_pages = True

    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self._pages = []

    def showPage(self):
        self._pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._pages)
        for state in self._pages:
            self.__dict__.update(state)
            w, h = self._pagesize
            self.setFont("Helvetica", 8)
            self.setFillColor(_c(MUTED))
            self.drawString(0.6 * inch, 0.45 * inch, self.header)
            if self.count_pages:
                self.drawRightString(w - 0.6 * inch, 0.45 * inch, f"Page {self._pageNumber} of {total}")
            self.setStrokeColor(_c(PURPLE))
            self.setLineWidth(3)
            self.line(0, h - 3, w, h - 3)
            super().showPage()
        super().save()


def _doc(buf, g: Grading, what: str, count_pages: bool = True):
    doc = SimpleDocTemplate(
        buf, pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
        topMargin=0.6 * inch, bottomMargin=0.7 * inch, title=f"{g.heading}: {what}",
    )

    class C(_NumberedCanvas):
        pass

    C.count_pages = count_pages
    C.header = " · ".join(x for x in (g.heading, g.course, what, _day(g.graded_at, "%b")) if x
    )

    return doc, C


def _table_style(header_rows: int = 1, zebra: bool = True) -> list:
    st = [
        ("BACKGROUND", (0, 0), (-1, header_rows - 1), _c(PURPLE)),
        ("TEXTCOLOR", (0, 0), (-1, header_rows - 1), colors.white),
        ("FONTNAME", (0, 0), (-1, header_rows - 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ("LINEBELOW", (0, header_rows), (-1, -1), 0.4, _c(LINE)),
    ]
    if zebra:
        st.append(("ROWBACKGROUNDS", (0, header_rows), (-1, -1), [colors.white, _c(ZEBRA)]))
    return st


def _tiles(pairs: list[tuple[str, str]], width: float) -> Table:
    cells = [[Paragraph(v, _TILE_NUM) for _, v in pairs], [Paragraph(k, _TILE_LBL) for k, _ in pairs]]
    t = Table(cells, colWidths=[width / len(pairs)] * len(pairs))
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), _c(ZEBRA)),
        ("LINEAFTER", (0, 0), (-2, -1), 3, colors.white),
        ("TOPPADDING", (0, 0), (-1, 0), 8),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
    ]))
    return t


def _bar(pct: float, w: float = 1.3 * inch, h: float = 7) -> Drawing:
    d = Drawing(w, h)
    d.add(Rect(0, 0, w, h, fillColor=_c(ZEBRA), strokeColor=None))
    d.add(Rect(0, 0, w * max(0, min(pct, 100)) / 100, h, fillColor=_c(PURPLE), strokeColor=None))
    return d


def _distribution_chart(g: Grading, width: float) -> Drawing:
    dist = g.distribution()
    d = Drawing(width, 150)
    ch = VerticalBarChart()
    ch.x, ch.y, ch.width, ch.height = 30, 25, width - 40, 110
    ch.data = [[n for _, n in dist]]
    ch.categoryAxis.categoryNames = [label for label, _ in dist]
    ch.categoryAxis.labels.fontName = ch.valueAxis.labels.fontName = "Helvetica"
    ch.categoryAxis.labels.fontSize = 7.5
    ch.categoryAxis.labels.fillColor = _c(MUTED)
    ch.categoryAxis.strokeColor = _c(LINE)
    top = max([n for _, n in dist] + [1])
    ch.valueAxis.valueMin, ch.valueAxis.valueMax = 0, top
    ch.valueAxis.valueStep = max(1, round(top / 4))
    ch.valueAxis.labels.fontSize = 7.5
    ch.valueAxis.labels.fillColor = _c(MUTED)
    ch.valueAxis.strokeColor = None
    ch.valueAxis.visibleGrid = True
    ch.valueAxis.gridStrokeColor = _c(LINE)
    ch.bars[0].fillColor = _c(PURPLE)
    ch.bars[0].strokeColor = None
    ch.barWidth, ch.groupSpacing = 10, 6
    d.add(ch)
    d.add(String(width / 2, 2, "Percent correct", fontName="Helvetica", fontSize=7.5, fillColor=_c(MUTED), textAnchor="middle"))
    return d


def to_pdf_report(g: Grading) -> bytes:
    """The class report: summary, roster, marks to check, item analysis."""
    buf = io.BytesIO()
    doc, canvasmaker = _doc(buf, g, "Class report")
    W = doc.width
    s = g.stats()
    story = [
        Paragraph(g.heading, _H1),
        Paragraph(" · ".join(x for x in (g.course, f"Graded {_stamp(g)}") if x), _SUB),
        _tiles([
            ("Students", str(s.get("count", 0))),
            ("Questions", str(len(g.questions))),
            ("Mean", f"{s.get('mean', 0):.1f}%"),
            ("Median", f"{s.get('median', 0):.1f}%"),
            ("Highest", f"{s.get('high', 0):.1f}%"),
            ("Lowest", f"{s.get('low', 0):.1f}%"),
        ], W),
        Spacer(1, 10),
        Paragraph("Score distribution", _H2),
        _distribution_chart(g, W),
    ]
    notes = g.key_warnings + g.skipped
    if notes:
        story += [Paragraph("Notes", _H2)] + [Paragraph("&bull; " + _esc(n), _BASE) for n in notes]

    key_line = "&nbsp;&nbsp; ".join(f"<font color='#{MUTED}'>{q}</font>&nbsp;<b>{g.key_answers[q]}</b>" for q in g.questions)
    story += [Paragraph("Answer key", _H2), Paragraph(key_line, ParagraphStyle("k", parent=_BASE, leading=15))]

    # Roster
    story += [PageBreak(), Paragraph("Scores", _H2)]
    rows = [["Name", "Sheet", "Score", "Percent", "", "To check"]]
    for r in g.by_name:
        name = Paragraph(_esc(r.sheet.name) if r.sheet.name else f"<i><font color='#{MUTED}'>No name (sheet {r.sheet.number})</font></i>", _CELL)
        rows.append([name, str(r.sheet.number), f"{r.score} / {r.possible}", f"{r.percent:.1f}%",
                     _bar(r.percent), str(len(r.flags)) if r.flags else ""])
    t = Table(rows, colWidths=[W * 0.36, W * 0.08, W * 0.12, W * 0.11, W * 0.21, W * 0.12], repeatRows=1)
    st = _table_style()
    st += [("ALIGN", (1, 0), (3, -1), "CENTER"), ("ALIGN", (5, 0), (5, -1), "CENTER"),
           ("TEXTCOLOR", (5, 1), (5, -1), _c(WARN)), ("FONTNAME", (5, 1), (5, -1), "Helvetica-Bold")]
    t.setStyle(TableStyle(st))
    story.append(t)

    # Marks to check
    flagged = [(r, q, f) for r in g.by_name for q, f in sorted(r.flags.items())]
    story.append(Paragraph("Marks to check on the paper sheet", _H2))
    if flagged:
        story.append(Paragraph(
            "<b>Multiple</b>: two bubbles filled, no credit given, as a machine would score it. "
            "<b>Light</b>: faint or an X, scored as read.", _SMALL))
        story.append(Spacer(1, 4))
        rows = [["Name", "Question", "Issue", "Read as", "Key"]]
        for r, q, f in flagged:
            rows.append([Paragraph(_esc(r.label), _CELL), str(q), f.capitalize(), _response(r, q), g.key_answers[q]])
        t = Table(rows, colWidths=[W * 0.4, W * 0.14, W * 0.16, W * 0.16, W * 0.14], repeatRows=1)
        st = _table_style() + [("ALIGN", (1, 0), (-1, -1), "CENTER"),
                               ("TEXTCOLOR", (2, 1), (2, -1), _c(WARN)), ("FONTNAME", (2, 1), (2, -1), "Helvetica-Bold")]
        t.setStyle(TableStyle(st))
        story.append(t)
    else:
        story.append(Paragraph(f"<font color='#{GOOD}'>Nothing to check: every mark was read cleanly.</font>", _BASE))

    # Item analysis
    story += [PageBreak(), Paragraph("Item analysis", _H2), Paragraph(
        "How many students chose each answer. The key's column is shaded. <b>Discrimination</b> is the "
        "share correct in the top third by score minus the bottom third: 0.3 or more is good, below 0 "
        "usually means a wrong key or a misleading question.", _SMALL), Spacer(1, 6)]
    head = ["Q", "Key", "% correct"] + list(CHOICES) + ["Blank", "Multi", "Discrim.", "Worth a look"]
    rows = [head]
    st = _table_style() + [("ALIGN", (0, 0), (-2, -1), "CENTER")]
    for i, it in enumerate(g.items, 1):
        d = it.discrimination
        rows.append([str(it.question), it.key, f"{it.percent:.0f}%"]
                    + [str(it.counts[c]) if it.counts[c] else "" for c in CHOICES]
                    + [str(it.counts[""] or ""), str(it.counts["*"] or ""),
                       "" if d is None else f"{d:.2f}", Paragraph(it.note, _SMALL) if it.note else ""])
        k = 3 + CHOICES.index(it.key)
        st += [("BACKGROUND", (k, i), (k, i), _c(GOOD_BG)), ("TEXTCOLOR", (k, i), (k, i), _c(GOOD)),
               ("FONTNAME", (k, i), (k, i), "Helvetica-Bold"), ("FONTNAME", (1, i), (1, i), "Helvetica-Bold"),
               ("TEXTCOLOR", (2, i), (2, i), _c(GOOD if it.percent >= 70 else WARN if it.percent >= 40 else BAD))]
        if d is not None and d < 0:
            st.append(("TEXTCOLOR", (10, i), (10, i), _c(BAD)))
    widths = [0.06, 0.06, 0.09] + [0.055] * 5 + [0.065, 0.065, 0.08] + [0.0]
    widths[-1] = 1 - sum(widths)
    t = Table(rows, colWidths=[W * x for x in widths], repeatRows=1)
    t.setStyle(TableStyle(st))
    story.append(t)

    doc.build(story, canvasmaker=canvasmaker)
    return buf.getvalue()


def to_pdf_students(g: Grading) -> bytes:
    """One page per student, alphabetical, to hand back or file."""
    buf = io.BytesIO()
    doc, canvasmaker = _doc(buf, g, "Student results", count_pages=False)
    W = doc.width
    mean = g.stats().get("mean", 0)
    story = []
    qs = g.questions
    per_col = 25
    cols_per_row = 4
    for n, r in enumerate(g.by_name):
        if n:
            story.append(PageBreak())
        story += [
            Paragraph(_esc(r.sheet.name) if r.sheet.name else f"No name (sheet {r.sheet.number})", _H1),
            Paragraph(" · ".join(x for x in (g.heading, g.course, _day(g.graded_at)) if x), _SUB),
            _tiles([("Score", f"{r.score} / {r.possible}"), ("Percent", f"{r.percent:.1f}%"),
                    ("Class mean", f"{mean:.1f}%"), ("Missed", str(r.possible - r.score))], W),
            Spacer(1, 12),
        ]
        chunks = [qs[i:i + per_col] for i in range(0, len(qs), per_col)]
        for band in range(0, len(chunks), cols_per_row):
            group = chunks[band:band + cols_per_row]
            story.append(_answer_grid(g, r, group, W, cols_per_row))
            story.append(Spacer(1, 10))
        legend = (f"<font color='#{GOOD}'>&#9632;</font> correct &nbsp;&nbsp; "
                  f"<font color='#{BAD}'>&#9632;</font> wrong (key shown) &nbsp;&nbsp; "
                  f"<font color='#{WARN}'>&#9632;</font> confirm on the paper sheet &nbsp;&nbsp; "
                  "- left blank &nbsp;&nbsp; Multiple: two bubbles filled")
        story.append(Paragraph(legend, _SMALL))
    if not story:
        story.append(Paragraph("No students were graded.", _BASE))
    doc.build(story, canvasmaker=canvasmaker)
    return buf.getvalue()


def _answer_grid(g: Grading, r: Result, group: list[list[int]], W: float, ncols: int) -> KeepTogether:
    """Up to four side-by-side columns of question / answer / key."""
    ncols_used = len(group)
    rows = [["Q", "Yours", "Key", ""] * ncols_used]
    height = max(len(c) for c in group)
    st = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("TEXTCOLOR", (0, 0), (-1, 0), _c(MUTED)),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
    for j in range(ncols_used):
        st.append(("LINEBELOW", (j * 4, 0), (j * 4 + 2, 0), 0.6, _c(PURPLE)))
    for i in range(height):
        row = []
        for j in range(ncols_used):
            col = group[j]
            if i >= len(col):
                row += ["", "", "", ""]
                continue
            q = col[i]
            yours = _response(r, q)
            row += [str(q), yours, "" if q in r.correct else g.key_answers[q], ""]
            c0 = j * 4
            if q in r.flags:
                bg, fg = WARN_BG, WARN
            elif q in r.correct:
                bg, fg = GOOD_BG, GOOD
            else:
                bg, fg = BAD_BG, BAD
            st += [("BACKGROUND", (c0, i + 1), (c0 + 2, i + 1), _c(bg)),
                   ("TEXTCOLOR", (c0 + 1, i + 1), (c0 + 1, i + 1), _c(fg)),
                   ("FONTNAME", (c0 + 1, i + 1), (c0 + 1, i + 1), "Helvetica-Bold"),
                   ("TEXTCOLOR", (c0, i + 1), (c0, i + 1), _c(MUTED)),
                   ("LINEBELOW", (c0, i + 1), (c0 + 2, i + 1), 0.8, colors.white)]
            if yours == "Multiple":
                st.append(("FONTSIZE", (c0 + 1, i + 1), (c0 + 1, i + 1), 6.5))
        rows.append(row)
    gap = 0.04
    unit = (1 - gap * (ncols - 1)) / ncols
    widths = []
    for j in range(ncols_used):
        widths += [W * unit * 0.28, W * unit * 0.40, W * unit * 0.32, W * gap if j < ncols_used - 1 else 0.001]
    t = Table(rows, colWidths=widths, hAlign="LEFT")
    t.setStyle(TableStyle(st))
    return KeepTogether(t)


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
