"""The sample PDF: a blank sheet, the key, and two students, both sides each.

Student B's sheet deliberately has a double mark (2), a blank (3), two X'd
bubbles (6, 9) and a half-erased second mark (7).
"""

import pathlib
import unittest

import io

import openpyxl
import pymupdf

from scantron import exports
from scantron.app import app
from scantron.grade import grade
from scantron.reader import read_pdf

SAMPLE = pathlib.Path(__file__).parent.parent / "samples" / "sample-key-and-students.pdf"


class SampleTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.sheets, cls.problems = read_pdf(SAMPLE.read_bytes(), "sample")

    def answers(self, i, n=10):
        a = self.sheets[i].answers
        return "".join(a[q].choice or "-" for q in range(1, n + 1))

    def test_pairs_sides_into_four_sheets(self):
        self.assertEqual(self.problems, [])
        self.assertEqual(len(self.sheets), 4)
        self.assertTrue(all(s.side1 and s.side2 for s in self.sheets))

    def test_blank_sheet(self):
        self.assertTrue(self.sheets[0].is_blank)

    def test_names(self):
        self.assertEqual([s.name for s in self.sheets[1:]], ["KEY", "STUDENT A", "STUDENT B"])

    def test_answers(self):
        self.assertEqual(self.answers(1), "ABCDEABCDE")
        self.assertEqual(self.answers(2), "ACCAEBBCAD")
        self.assertEqual(self.answers(3)[2], "-")

    def test_nothing_read_past_question_10(self):
        for s in self.sheets:
            self.assertFalse([q for q, a in s.answers.items() if q > 10 and a.choice])

    def test_flags(self):
        b = self.sheets[3].answers
        self.assertEqual(b[2].flag, "multiple")
        self.assertEqual((b[6].choice, b[6].flag), ("B", "light"))
        self.assertEqual((b[9].choice, b[9].flag), ("D", "light"))
        self.assertEqual((b[7].choice, b[7].flag), ("B", ""))

    def test_scores(self):
        g = grade(self.sheets[1], self.sheets[2:], [])
        self.assertEqual([r.score for r in g.results], [5, 6])
        self.assertEqual(g.results[1].responses[2], "*")


class ExportTest(unittest.TestCase):
    """Every download builds, and carries the same scores as the page."""

    @classmethod
    def setUpClass(cls):
        c = app.test_client()
        r = c.post(
            "/grade",
            data={"students": (open(SAMPLE, "rb"), "sample.pdf"), "title": "Quiz 1", "course": "TEST 101"},
            content_type="multipart/form-data",
        )
        cls.client, cls.base = c, r.location

    def get(self, kind):
        r = self.client.get(f"{self.base}/download/{kind}")
        self.assertEqual(r.status_code, 200, kind)
        self.assertIn("TEST 101 Quiz 1", r.headers["Content-Disposition"])
        return r.data

    def test_csv(self):
        rows = self.get("csv").decode("utf-8-sig").splitlines()
        self.assertTrue(rows[1].startswith("ANSWER KEY"))
        self.assertTrue(rows[2].startswith("STUDENT A,1,5,10,50.0"))
        self.assertTrue(rows[3].startswith("STUDENT B,2,6,10,60.0,3"))

    def test_xlsx(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.get("xlsx")))
        self.assertEqual(wb.sheetnames, ["Summary", "Scores", "Responses", "Item Analysis", "To Check"])
        scores = wb["Scores"]
        self.assertEqual([scores.cell(row=r, column=3).value for r in (5, 6)], [5, 6])
        self.assertEqual(wb["To Check"].max_row, 7)  # header at 4, three marks

    def test_pdfs(self):
        report = pymupdf.open(stream=self.get("report.pdf"), filetype="pdf")
        text = "".join(p.get_text() for p in report)
        self.assertIn("Quiz 1", text)
        self.assertIn("STUDENT B", text)
        students = pymupdf.open(stream=self.get("students.pdf"), filetype="pdf")
        self.assertEqual(len(students), 2)
        self.assertIn("5 / 10", students[0].get_text())


if __name__ == "__main__":
    unittest.main()
